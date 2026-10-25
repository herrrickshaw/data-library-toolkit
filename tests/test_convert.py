"""Tests for dlt/convert.py and dlt/analyze.py with generated fixtures.
rclone is a stub script copying between temp folders; nothing leaves tmp_path."""

import stat
import sys
import textwrap
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from dlt import analyze, catalog, config, convert  # noqa: E402


def make_pdf(path, text="Hello DLT"):
    import fitz
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_paragraph("Para one")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "c1", "c2"
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


# ---- pure helpers --------------------------------------------------------------

@pytest.mark.parametrize("title,expected", [
    ("Paper.PDF", "Paper"), ("x/y?.epub", "x_y_"), (None, "untitled"), ("v1.2 notes", "v1.2 notes"),
])
def test_safe_title_for(title, expected):
    assert convert.safe_title_for(title) == expected


def test_safe_dirname_and_ext():
    assert convert._safe_dirname("s3:bucket/x") == "s3_bucket_x"
    assert convert.ext_for("a/b.Pdf") == ".Pdf"
    assert convert.ext_for("a/b") == ".bin"


def test_extract_docx_and_corrupt(tmp_path):
    make_docx(tmp_path / "a.docx")
    assert convert.extract_text(tmp_path / "a.docx", "docx") == ("Para one\n\nc1 | c2", None)
    (tmp_path / "bad.pptx").write_bytes(b"nope")
    text, err = convert.extract_text(tmp_path / "bad.pptx", "pptx")
    assert text is None and err.startswith("pptx extraction error")


def test_extract_pdf(tmp_path):
    make_pdf(tmp_path / "a.pdf")
    text, err = convert.extract_text(tmp_path / "a.pdf", "pdf")
    assert err is None and "Hello DLT" in text


# ---- run() end-to-end ---------------------------------------------------------

@pytest.fixture
def env(tmp_path, monkeypatch):
    cloud, uploaded = tmp_path / "cloud", tmp_path / "uploaded"
    cloud.mkdir()
    stub = tmp_path / "rclone"
    stub.write_text(textwrap.dedent(f"""\
        #!/bin/sh
        src="$2"; dst="$3"; list=""; prev=""
        for a in "$@"; do [ "$prev" = "--files-from" ] && list="$a"; prev="$a"; done
        if [ -n "$list" ]; then
          while IFS= read -r p || [ -n "$p" ]; do
            [ -f "{cloud}/$p" ] || continue
            mkdir -p "$dst/$(dirname "$p")"; cp "{cloud}/$p" "$dst/$p"
          done < "$list"
        else
          mkdir -p "{uploaded}/$dst"; cp -R "$src"/. "{uploaded}/$dst"
        fi
        """))
    stub.chmod(stub.stat().st_mode | stat.S_IEXEC)
    data = tmp_path / "data"
    monkeypatch.setattr(config, "RCLONE_BIN", str(stub))
    monkeypatch.setattr(config, "DATA_DIR", data)
    monkeypatch.setattr(config, "CATALOG_DB", data / "c.duckdb")
    monkeypatch.setattr(config, "MD_OUTPUT_DIR", data / "md")
    monkeypatch.setattr(config, "STAGE_DIR", data / "stage")
    monkeypatch.setattr(config, "FETCH_DIR", data / "stage" / "fetch")
    monkeypatch.setattr(config, "UPLOAD_DIR", data / "stage" / "upload")
    monkeypatch.setattr(config, "DISK_SAFETY_MARGIN_GB", 0)
    monkeypatch.setattr(config, "CLOUD_DEST_FOLDER", "out")
    config.ensure_dirs()
    return type("E", (), {"cloud": cloud, "uploaded": uploaded, "data": data})


def _add(key, path, filetype="pdf", source="dropbox:", title=None):
    con = catalog.connect()
    con.execute("insert into documents (item_key, title, path, source, filetype) values (?,?,?,?,?)",
                [key, title or Path(path).name, path, source, filetype])
    con.close()


def _status():
    con = catalog.connect()
    out = {k: (s, c) for k, s, c in con.execute("select item_key, status, cloud_status from fulltext_md").fetchall()}
    con.close()
    return out


def test_run_converts_uploads_and_classifies(env):
    make_pdf(env.cloud / "p" / "a.pdf")
    make_docx(env.cloud / "p" / "b.docx")
    _add("K1", "p/a.pdf")
    _add("K2", "p/b.docx", "docx")
    _add("K3", "p/z.zip!in.pdf")
    _add("K4", "p/missing.pdf")
    res = convert.run(max_docs=10, log=lambda *a: None)
    assert _status() == {"K1": ("ok", "ok"), "K2": ("ok", "ok"),
                         "K3": ("archive_not_persisted", None), "K4": ("fetch_failed", None)}
    assert res == {"converted": 2, "remaining": 0}
    up = sorted(p.name for p in env.uploaded.rglob("*") if p.is_file())
    assert up == ["K1_a.md", "K1_a.pdf", "K2_b.docx", "K2_b.md"]


def test_run_respects_max_docs_and_resumes(env):
    for i in range(3):
        make_pdf(env.cloud / f"d{i}.pdf")
        _add(f"M{i}", f"d{i}.pdf")
    assert convert.run(max_docs=2, log=lambda *a: None)["remaining"] == 1
    assert convert.run(max_docs=2, log=lambda *a: None)["remaining"] == 0
    assert set(_status()) == {"M0", "M1", "M2"}


def test_status_only_and_analyze_report(env):
    make_pdf(env.cloud / "a.pdf")
    _add("S1", "a.pdf")
    convert.run(max_docs=5, log=lambda *a: None)
    s = convert.run(status_only=True, log=lambda *a: None)
    assert s["total"] == 1 and s["remaining"] == 0
    con = catalog.connect()
    out = analyze.generate(con, out_path=env.data / "A.md", log=lambda *a: None)
    txt = out.read_text()
    assert "| Catalogued documents | 1 |" in txt and "| ok | 1 |" in txt
